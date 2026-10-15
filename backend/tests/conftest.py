from io import BytesIO

import fitz
import pytest
from docx import Document


@pytest.fixture
def sample_pdf_bytes() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        "Jane Developer\n\nExperience\n• Built APIs with FastAPI\n• Led migration to cloud",
        fontsize=12,
    )
    payload = document.tobytes()
    document.close()
    return payload


@pytest.fixture
def sample_docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Jane Developer", level=1)
    document.add_paragraph("Professional summary with more than forty characters of text.")
    document.add_paragraph("Built APIs with FastAPI", style="List Bullet")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def scanned_pdf_bytes() -> bytes:
    document = fitz.open()
    document.new_page()
    payload = document.tobytes()
    document.close()
    return payload
